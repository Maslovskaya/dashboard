"""Создаёт data/cpu.xlsx. Цены и характеристики ориентировочные:
перед защитой замените их на актуальные данные из магазина."""
from pathlib import Path
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# (модель, производитель, сокет, ядра, потоки, частота ГГц, TDP Вт, цена руб)
CPUS = [
    ("Ryzen 3 4100", "AMD", "AM4", 4, 8, 4.0, 65, 5500),
    ("Ryzen 5 5500", "AMD", "AM4", 6, 12, 4.2, 65, 7500),
    ("Ryzen 5 5600", "AMD", "AM4", 6, 12, 4.4, 65, 9500),
    ("Ryzen 5 5600X", "AMD", "AM4", 6, 12, 4.6, 65, 11500),
    ("Ryzen 7 5700X", "AMD", "AM4", 8, 16, 4.6, 65, 15500),
    ("Ryzen 7 5800X3D", "AMD", "AM4", 8, 16, 4.5, 105, 27000),
    ("Ryzen 9 5900X", "AMD", "AM4", 12, 24, 4.8, 105, 24000),
    ("Ryzen 9 5950X", "AMD", "AM4", 16, 32, 4.9, 105, 32000),
    ("Ryzen 5 7600", "AMD", "AM5", 6, 12, 5.1, 65, 17500),
    ("Ryzen 5 7600X", "AMD", "AM5", 6, 12, 5.3, 105, 19500),
    ("Ryzen 7 7700", "AMD", "AM5", 8, 16, 5.3, 65, 24500),
    ("Ryzen 7 7700X", "AMD", "AM5", 8, 16, 5.4, 105, 26500),
    ("Ryzen 7 7800X3D", "AMD", "AM5", 8, 16, 5.0, 120, 36000),
    ("Ryzen 9 7900", "AMD", "AM5", 12, 24, 5.4, 65, 34000),
    ("Ryzen 9 7900X", "AMD", "AM5", 12, 24, 5.6, 170, 36500),
    ("Ryzen 9 7950X", "AMD", "AM5", 16, 32, 5.7, 170, 48000),
    ("Ryzen 5 9600X", "AMD", "AM5", 6, 12, 5.4, 65, 22500),
    ("Ryzen 7 9700X", "AMD", "AM5", 8, 16, 5.5, 65, 32000),
    ("Ryzen 7 9800X3D", "AMD", "AM5", 8, 16, 5.2, 120, 46000),
    ("Ryzen 9 9950X", "AMD", "AM5", 16, 32, 5.7, 170, 57000),
    ("Core i3-12100F", "Intel", "LGA1700", 4, 8, 4.3, 58, 6500),
    ("Core i3-14100F", "Intel", "LGA1700", 4, 8, 4.7, 58, 8000),
    ("Core i5-12400F", "Intel", "LGA1700", 6, 12, 4.4, 65, 10000),
    ("Core i5-12600KF", "Intel", "LGA1700", 10, 16, 4.9, 125, 16000),
    ("Core i7-12700KF", "Intel", "LGA1700", 12, 20, 5.0, 125, 22500),
    ("Core i5-13400F", "Intel", "LGA1700", 10, 16, 4.6, 65, 15500),
    ("Core i5-13600KF", "Intel", "LGA1700", 14, 20, 5.1, 125, 21500),
    ("Core i7-13700KF", "Intel", "LGA1700", 16, 24, 5.4, 125, 30000),
    ("Core i9-13900KF", "Intel", "LGA1700", 24, 32, 5.8, 125, 42000),
    ("Core i5-14400F", "Intel", "LGA1700", 10, 16, 4.7, 65, 16500),
    ("Core i5-14600KF", "Intel", "LGA1700", 14, 20, 5.3, 125, 24000),
    ("Core i7-14700KF", "Intel", "LGA1700", 20, 28, 5.6, 125, 33500),
    ("Core i9-14900KF", "Intel", "LGA1700", 24, 32, 6.0, 125, 45000),
    ("Core Ultra 5 245K", "Intel", "LGA1851", 14, 14, 5.2, 125, 26500),
    ("Core Ultra 7 265K", "Intel", "LGA1851", 20, 20, 5.5, 125, 36500),
    ("Core Ultra 9 285K", "Intel", "LGA1851", 24, 24, 5.7, 125, 58000),
]
COLUMNS = ["Модель", "Производитель", "Сокет", "Ядра", "Потоки",
           "Частота (ГГц)", "TDP (Вт)", "Цена (руб)"]

out = Path(__file__).parent / "cpu.xlsx"
df = pd.DataFrame(CPUS, columns=COLUMNS)
with pd.ExcelWriter(out, engine="openpyxl") as w:
    df.to_excel(w, index=False, sheet_name="Процессоры")
    ws = w.sheets["Процессоры"]
    head_fill = PatternFill("solid", fgColor="2F5D8A")
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial")
    for i, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(col) + 4)
    ws.column_dimensions["A"].width = 22
    ws.freeze_panes = "A2"
print("OK", out, len(df), "строк")
